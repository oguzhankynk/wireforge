from __future__ import annotations

import base64
import csv
import importlib
import io
import shutil
from pathlib import Path

import openpyxl
import pytest
from fastapi.testclient import TestClient

from wireforge import cli, importers, service
from wireforge.codec import ReferenceCodec
from wireforge.golden import Sampler
from wireforge.loader import load_project
from wireforge.server.app import create_app

ICD = Path(__file__).resolve().parents[2] / "examples" / "icd"
CSV = ICD / "telemetry_icd.csv"
HEADER = ICD / "sensors.h"
OPTIONS = importers.Project("telemetry", "demo::imu", endian="little")


def _round_trip(location: Path) -> None:
    project = load_project(location).project
    assert project is not None
    codec = ReferenceCodec(project)
    for message in project.messages.values():
        value = Sampler(codec).struct(message.fields)
        payload = codec.encode_message(message.name, value)
        assert (
            codec.encode_message(message.name, codec.decode_message(message.name, payload))
            == payload
        )


def test_sample_csv_imports_to_a_valid_project(tmp_path: Path) -> None:
    result = service.import_icd(CSV, tmp_path / "telemetry", OPTIONS)
    assert result.ok, result.text
    text = (tmp_path / "telemetry" / "project.yaml").read_text(encoding="utf-8")
    assert "    id: 0x0110\n" in text
    assert "        occurs: { max: 8, count: sampleCount }\n" in text
    assert (
        "      - { name: mode, type: u8, bits: 3, min: 0, max: 5, description: Operating mode }"
        in text
    )
    assert service.format_project(tmp_path / "telemetry", check=True).ok
    assert service.generate_code(tmp_path / "telemetry", tmp_path / "out").ok
    _round_trip(tmp_path / "telemetry")


def test_excel_matches_csv(tmp_path: Path) -> None:
    workbook = openpyxl.Workbook()
    sheet = workbook.active
    assert sheet is not None
    for row in csv.reader(io.StringIO(CSV.read_text(encoding="utf-8"))):
        sheet.append([float(cell) if cell.lstrip("-").isdigit() else cell for cell in row])
    path = tmp_path / "icd.xlsx"
    workbook.save(path)
    assert importers.from_xlsx(path, OPTIONS) == importers.from_csv(
        CSV.read_text(encoding="utf-8"), OPTIONS
    )
    empty = tmp_path / "empty.xlsx"
    openpyxl.Workbook().save(empty)
    with pytest.raises(importers.IcdImportError, match="empty"):
        importers.from_xlsx(empty, OPTIONS)


def test_excel_needs_openpyxl(tmp_path: Path, monkeypatch: pytest.MonkeyPatch) -> None:
    real = importlib.import_module

    def missing(name: str) -> object:
        if name == "openpyxl":
            raise ImportError(name)
        return real(name)

    monkeypatch.setattr(importlib, "import_module", missing)
    with pytest.raises(importers.IcdImportError, match="openpyxl"):
        importers.from_xlsx(tmp_path / "x.xlsx", OPTIONS)


def test_sample_header_imports_to_a_valid_project(tmp_path: Path) -> None:
    result = service.import_icd(HEADER, tmp_path / "sensors", importers.Project("sensors"))
    assert result.ok, result.text
    project = load_project(tmp_path / "sensors").project
    assert project is not None
    assert list(project.messages) == ["SensorFrame", "Ping"]
    assert list(project.types) == ["Flags"]
    assert project.enums["Mode"].values == {"MODE_IDLE": 0, "MODE_RUN": 1, "MODE_FAULT": 7}
    _round_trip(tmp_path / "sensors")
    chosen = importers.from_c_header(
        HEADER.read_text(encoding="utf-8"), importers.Project("s"), ["Ping"]
    )
    assert list(chosen["messages"]) == ["Ping"]
    assert set(chosen["types"]) == {"Flags", "SensorFrame"}


def test_header_enum_widths_and_errors() -> None:
    options = importers.Project("p")
    wide = importers.from_c_header(
        "enum Big { A = 0x1FFFF };\nenum Mid { B = 300 };\nenum Neg { C = -1, D };\n", options, []
    )
    assert {name: enum["type"] for name, enum in wide["enums"].items()} == {
        "Big": "u32",
        "Mid": "u16",
        "Neg": "i32",
    }
    assert wide["enums"]["Neg"]["values"] == {"C": -1, "D": 0}
    cases = {
        "#define GAIN 1.5f\n": "only integer defines",
        "struct S {\n  uint8_t *p;\n};\n": ":2: /: pointers and unions",
        "struct S { long x; };": "unknown type 'long'",
        "struct S { uint8_t a b c; };": "cannot read member",
        "enum E { A = X };": "is not an integer",
    }
    for text, message in cases.items():
        with pytest.raises(importers.IcdImportError, match=message):
            importers.from_c_header(text, options, source="h")
    with pytest.raises(importers.IcdImportError, match="no struct 'Missing'"):
        importers.from_c_header("struct S { uint8_t a; };", options, ["Missing"])


def test_csv_errors() -> None:
    options = importers.Project("p")
    cases = {
        "name,type\nM,u8\n": "missing columns: field",
        "name,field,type,colour\nM,a,u8,red\n": "unknown columns: colour",
        "section,name,field,type\nthing,M,a,u8\n": "section must be",
        "name,field,type\nM,,u8\n": "needs name, field and type",
        "name,field,type,bits\nM,a,u8,x\n": "invalid literal",
    }
    for text, message in cases.items():
        with pytest.raises(importers.IcdImportError, match=message):
            importers.from_csv(text, options)
    document = importers.from_csv(
        "name,field,type,occurs,min,max,default\nM,a,u8,fill..16,,,\n,,,,,,\nN,b,f32,1..4,-1.5,2e3,0.25\n",
        options,
    )
    assert document["messages"]["M"]["fields"] == [
        {"name": "a", "type": "u8", "occurs": {"max": 16, "fill": True}}
    ]
    assert document["messages"]["N"]["fields"][0] == {
        "name": "b",
        "type": "f32",
        "min": -1.5,
        "max": 2000.0,
        "default": 0.25,
        "occurs": {"max": 4, "min": 1},
    }


def test_service_and_cli(tmp_path: Path) -> None:
    occupied = tmp_path / "occupied"
    occupied.mkdir()
    (occupied / "x").write_text("x", encoding="utf-8")
    with pytest.raises(service.ServiceError, match="not empty"):
        service.import_icd(CSV, occupied, OPTIONS)
    with pytest.raises(service.ServiceError, match=r"reads \.csv"):
        service.import_icd(Path("icd.txt"), tmp_path / "t", OPTIONS)
    bad = tmp_path / "bad.csv"
    bad.write_text("name,type\n", encoding="utf-8")
    assert not service.import_icd(bad, tmp_path / "bad", OPTIONS).ok
    target = tmp_path / "cli"
    code = cli.main(
        ["import-icd", str(HEADER), str(target), "--name", "sensors", "--message", "Ping"]
    )
    assert code == cli.EXIT_OK
    assert cli.main(["import-icd", str(HEADER), str(target), "--name", "sensors"]) == cli.EXIT_ERROR


def test_server_import(tmp_path: Path) -> None:
    client = TestClient(create_app(tmp_path))
    csv_text = CSV.read_text(encoding="utf-8")
    created = client.post(
        "/api/import", json={"id": "imu", "name": "telemetry", "format": "csv", "content": csv_text}
    )
    assert created.status_code == 201
    assert created.json()["ok"]
    assert (
        client.post(
            "/api/import", json={"id": "imu", "name": "x", "format": "csv", "content": csv_text}
        ).status_code
        == 409
    )
    header = client.post(
        "/api/import",
        json={
            "id": "sensors",
            "name": "sensors",
            "format": "h",
            "content": HEADER.read_text(encoding="utf-8"),
        },
    )
    assert header.json()["ok"]
    workbook = openpyxl.Workbook()
    sheet = workbook.active
    assert sheet is not None
    for row in csv.reader(io.StringIO(csv_text)):
        sheet.append(row)
    buffer = io.BytesIO()
    workbook.save(buffer)
    excel = client.post(
        "/api/import",
        json={
            "id": "excel",
            "name": "telemetry",
            "format": "xlsx",
            "content": base64.b64encode(buffer.getvalue()).decode(),
        },
    )
    assert excel.json()["ok"]
    assert (
        client.post(
            "/api/import", json={"id": "x", "name": "x", "format": "xlsx", "content": "%%%"}
        ).status_code
        == 400
    )
    failed = client.post(
        "/api/import", json={"id": "broken", "name": "x", "format": "csv", "content": "name,type\n"}
    )
    assert not failed.json()["ok"]
    assert not (tmp_path / "broken").exists()
    assert [item["id"] for item in client.get("/api/projects").json()] == [
        "excel",
        "imu",
        "sensors",
    ]
    shutil.rmtree(tmp_path / "excel")
